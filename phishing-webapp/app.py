from flask import (
    Flask, request, jsonify, send_from_directory,
    redirect, session
)
import joblib, pandas as pd, numpy as np
import re, math, string, os, base64, io, json, traceback, zipfile, struct

from google_auth_oauthlib.flow import Flow
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
import google.auth.transport.requests
import urllib.request
import difflib
from email.utils import parseaddr

# ── OCR: Pillow + Tesseract ──────────────────────────────────────
try:
    from PIL import Image
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

# ── PDF ──────────────────────────────────────────────────────────
try:
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# ── DOCX ─────────────────────────────────────────────────────────
try:
    import docx as python_docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# ── XLSX ─────────────────────────────────────────────────────────
try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

# ── HTML parsing ─────────────────────────────────────────────────
try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False

# ── oletools for macro analysis ──────────────────────────────────
try:
    from oletools.olevba import VBA_Parser
    OLETOOLS_AVAILABLE = True
except ImportError:
    OLETOOLS_AVAILABLE = False

# ── pefile for EXE analysis ──────────────────────────────────────
try:
    import pefile
    PEFILE_AVAILABLE = True
except ImportError:
    PEFILE_AVAILABLE = False

os.environ["OAUTHLIB_INSECURE_TRANSPORT"] = "1"

app = Flask(__name__, static_folder="static")
app.config['SESSION_COOKIE_SAMESITE'] = "Lax"
app.config['SESSION_COOKIE_SECURE']   = False
app.secret_key = os.environ.get("FLASK_SECRET", "CHANGE_ME_IN_PRODUCTION_abc123xyz")

BASE   = os.path.dirname(os.path.abspath(__file__))
MODELS = os.path.join(BASE, "models")

# ════════════════════════════════════════════════════════════════
#  GEMINI CONFIG
# ════════════════════════════════════════════════════════════════

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "AIzaSyBqAS17MfHWaSNC1uSvNlE89YSlLk_7EHA")
GEMINI_API_URL = "https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent"

def call_gemini(prompt: str) -> str:
    """Call Gemini API and return text response."""
    try:
        payload = json.dumps({
            "contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {
                "temperature": 0.3,
                "maxOutputTokens": 1024,
            }
        }).encode("utf-8")

        req = urllib.request.Request(
            f"{GEMINI_API_URL}?key={GEMINI_API_KEY}",
            data=payload,
            headers={"Content-Type": "application/json"},
            method="POST"
        )
        with urllib.request.urlopen(req, timeout=15) as resp:
            result = json.loads(resp.read().decode())
            return result["candidates"][0]["content"]["parts"][0]["text"]
    except Exception as e:
        print(f"❌ Gemini API error: {e}")
        return None


# ════════════════════════════════════════════════════════════════
#  LOAD MODELS
# ════════════════════════════════════════════════════════════════

def load_models():
    try:
        um  = joblib.load(os.path.join(MODELS, "url_model.pkl"))
        us  = joblib.load(os.path.join(MODELS, "url_scaler.pkl"))
        tm  = joblib.load(os.path.join(MODELS, "text_model.pkl"))
        tf_ = joblib.load(os.path.join(MODELS, "tfidf.pkl"))
        print("✅ All 4 models loaded successfully.")
        return um, us, tm, tf_
    except FileNotFoundError as e:
        print(f"⚠️  Model files not found: {e}")
        return None, None, None, None

url_model, url_scaler, text_model, tfidf = load_models()

# ════════════════════════════════════════════════════════════════
#  OAUTH CONFIG
# ════════════════════════════════════════════════════════════════

CLIENT_SECRETS_FILE = os.path.join(BASE, "client_secret.json")
SCOPES = [
    "openid",
    "https://www.googleapis.com/auth/userinfo.email",
    "https://www.googleapis.com/auth/userinfo.profile",
    "https://www.googleapis.com/auth/gmail.readonly",
]
REDIRECT_URI = "http://localhost:5000/oauth2callback"

# ════════════════════════════════════════════════════════════════
#  HELPERS — ORIGINAL
# ════════════════════════════════════════════════════════════════

def get_credentials():
    creds_data = session.get("credentials")
    if not creds_data:
        return None
    creds = Credentials(
        token         = creds_data["token"],
        refresh_token = creds_data.get("refresh_token"),
        token_uri     = creds_data["token_uri"],
        client_id     = creds_data["client_id"],
        client_secret = creds_data["client_secret"],
        scopes        = creds_data["scopes"],
    )
    if creds.expired and creds.refresh_token:
        try:
            creds.refresh(google.auth.transport.requests.Request())
            session["credentials"]["token"] = creds.token
            print("🔄 Token refreshed.")
        except Exception as e:
            print(f"❌ Token refresh failed: {e}")
            return None
    return creds


def get_trusted_contacts(service) -> set:
    trusted = set()
    print("🔍 Fetching trusted contacts from SENT folder...")
    try:
        resp = service.users().messages().list(userId="me", maxResults=20, labelIds=["SENT"]).execute()
        msgs = resp.get("messages", [])
        for msg_meta in msgs:
            try:
                msg_data = service.users().messages().get(userId="me", id=msg_meta["id"], format="metadata", metadataHeaders=["To"]).execute()
                headers = msg_data.get("payload", {}).get("headers", [])
                for h in headers:
                    if h["name"].lower() == "to":
                        emails = re.findall(r'[\w\.-]+@[\w\.-]+\.\w+', h["value"])
                        for e in emails:
                            trusted.add(e.lower())
            except Exception:
                pass
        print(f"✅ Found {len(trusted)} trusted contacts.")
    except Exception as e:
        print(f"⚠️ Failed to fetch trusted contacts: {e}")
    return trusted


def check_impersonation(sender: str, trusted_contacts: set) -> dict:
    _, email_addr = parseaddr(sender)
    email_addr = email_addr.lower()

    if not email_addr or "@" not in email_addr:
        return {"is_impersonation": False, "score": 0, "matched_contact": None}

    username, domain = email_addr.split("@", 1)

    best_score = 0
    best_contact = None

    for contact in trusted_contacts:
        if "@" not in contact: continue
        c_user, c_domain = contact.split("@", 1)

        if domain == c_domain and username == c_user:
            continue

        score = difflib.SequenceMatcher(None, username, c_user).ratio()
        if score > best_score:
            best_score = score
            best_contact = contact

    if best_score > 0.85:
        return {"is_impersonation": True, "score": round(best_score * 100, 1), "matched_contact": best_contact}

    return {"is_impersonation": False, "score": round(best_score * 100, 1), "matched_contact": None}


def extract_url_features(url: str) -> list:
    digits = sum(c.isdigit() for c in url)
    return [
        len(url),
        url.count("."),
        url.count("-"),
        url.count("/"),
        url.count("@"),
        url.count("?"),
        url.count("="),
        1 if re.search(r"\d+\.\d+\.\d+\.\d+", url) else 0,
        1 if url.lower().startswith("https") else 0,
        round(digits / max(len(url), 1), 4),
    ]


def get_url_red_flags(url: str) -> list:
    flags = []
    if re.search(r"\d+\.\d+\.\d+\.\d+", url):
        flags.append("Uses raw IP address instead of domain name")
    if not url.lower().startswith("https"):
        flags.append("Not using HTTPS (insecure connection)")
    if url.count("-") > 3:
        flags.append(f"Excessive hyphens ({url.count('-')}) in URL — common phishing tactic")
    if len(url) > 75:
        flags.append(f"Unusually long URL ({len(url)} characters)")
    if url.count(".") > 4:
        flags.append(f"Too many subdomains ({url.count('.')} dots)")
    if "@" in url:
        flags.append("Contains @ symbol — hides the real destination")
    suspicious_words = ["login", "verify", "secure", "update", "confirm", "account", "bank", "paypal", "amazon"]
    found = [w for w in suspicious_words if w in url.lower()]
    if found:
        flags.append(f"Contains urgency/trust keywords: {', '.join(found)}")
    digit_ratio = sum(c.isdigit() for c in url) / max(len(url), 1)
    if digit_ratio > 0.2:
        flags.append(f"High digit ratio ({round(digit_ratio*100)}%) — typical of auto-generated phishing domains")
    return flags


def decode_body(payload):
    if "parts" in payload:
        for part in payload["parts"]:
            if part.get("mimeType") == "text/plain":
                data = part.get("body", {}).get("data", "")
                if data:
                    return base64.urlsafe_b64decode(data + "==").decode("utf-8", errors="ignore")
        for part in payload["parts"]:
            result = decode_body(part)
            if result:
                return result
    else:
        data = payload.get("body", {}).get("data", "")
        if data:
            return base64.urlsafe_b64decode(data + "==").decode("utf-8", errors="ignore")
    return ""


def run_text_model(text: str) -> dict:
    """Run text through tfidf + text_model. Returns prediction and confidence."""
    if text_model is None or tfidf is None:
        return {"is_phishing": False, "confidence": 0.0}
    try:
        vec   = tfidf.transform([text])
        pred  = int(text_model.predict(vec)[0])
        proba = text_model.predict_proba(vec)[0]
        return {"is_phishing": bool(pred == 1), "confidence": round(float(max(proba)) * 100, 2)}
    except Exception:
        return {"is_phishing": False, "confidence": 0.0}


def run_url_model(url: str) -> bool:
    """Run a URL through the url_model. Returns True if phishing."""
    if url_model is None or url_scaler is None:
        return False
    try:
        f      = extract_url_features(url)
        scaled = url_scaler.transform([f])
        return int(url_model.predict(scaled)[0]) == 1
    except Exception:
        return False


# ════════════════════════════════════════════════════════════════
#  ATTACHMENT ANALYSIS HELPERS
# ════════════════════════════════════════════════════════════════

# ── Magic bytes for file type verification ───────────────────────
MAGIC_BYTES = {
    "exe":  b"MZ",
    "pdf":  b"%PDF",
    "png":  b"\x89PNG",
    "jpg":  b"\xff\xd8\xff",
    "zip":  b"PK\x03\x04",
    "gif":  b"GIF8",
    "webp": b"RIFF",
}

# ── Extensions considered dangerous without further analysis ─────
INSTANT_DANGER_EXTENSIONS = {".bat", ".cmd", ".ps1", ".vbs", ".js", ".hta", ".scr", ".pif"}

# ── Macro-enabled Office formats ─────────────────────────────────
MACRO_EXTENSIONS = {".xlsm", ".docm", ".xltm", ".dotm", ".pptm"}

# ── Suspicious EXE import names ──────────────────────────────────
SUSPICIOUS_IMPORTS = {
    "VirtualAlloc", "WriteProcessMemory", "CreateRemoteThread",
    "URLDownloadToFile", "WinExec", "ShellExecute", "CreateProcess",
    "RegSetValue", "IsDebuggerPresent",
}

# ── Suspicious EXE strings ───────────────────────────────────────
SUSPICIOUS_EXE_STRINGS = [
    "powershell", "cmd.exe", "wget", "curl", "reg add",
    "HKEY_", "taskkill", "netstat", "base64", "http://", "https://",
]

# ── Suspicious macro keywords ────────────────────────────────────
SUSPICIOUS_MACRO_KEYWORDS = [
    "Shell", "CreateObject", "WScript", "URLDownloadToFile",
    "Auto_Open", "Document_Open", "Workbook_Open", "AutoOpen",
    "AutoExec", "powershell", "cmd", "reg add", "VirtualAlloc",
]

# ── JS obfuscation markers (HTML) ────────────────────────────────
OBFUSCATION_MARKERS = ["eval(", "atob(", "unescape(", "String.fromCharCode(", "document.write("]


def calculate_entropy(data: bytes) -> float:
    """Calculate Shannon entropy of bytes. Scale 0–8. Above 7.2 = suspicious."""
    if not data:
        return 0.0
    freq = {}
    for byte in data:
        freq[byte] = freq.get(byte, 0) + 1
    entropy = 0.0
    length  = len(data)
    for count in freq.values():
        p = count / length
        if p > 0:
            entropy -= p * math.log2(p)
    return round(entropy, 3)


def verify_magic_bytes(data: bytes, claimed_ext: str) -> bool:
    """Returns True if file bytes match the claimed extension."""
    ext = claimed_ext.lower().lstrip(".")
    magic = MAGIC_BYTES.get(ext)
    if magic is None:
        return True   # Unknown type — can't verify, assume ok
    return data[:len(magic)] == magic


def extract_printable_strings(data: bytes, min_len: int = 6) -> list:
    """Extract printable ASCII strings from binary data."""
    result = []
    current = []
    for byte in data:
        if 32 <= byte < 127:
            current.append(chr(byte))
        else:
            if len(current) >= min_len:
                result.append("".join(current))
            current = []
    if len(current) >= min_len:
        result.append("".join(current))
    return result


# ── Per-type analyzers ───────────────────────────────────────────

def analyze_exe(data: bytes, filename: str) -> dict:
    """
    Static analysis of EXE/binary files.
    Never executes — read-only forensic analysis.
    """
    flags      = []
    risk_score = 0

    # 1. Magic byte check
    ext = os.path.splitext(filename)[1].lower()
    if not verify_magic_bytes(data, "exe"):
        flags.append(f"File extension is '{ext}' but internal bytes don't match a real EXE — possible disguised file")
        risk_score += 30

    # 2. Entropy
    entropy = calculate_entropy(data)
    flags.append(f"File entropy: {entropy}/8.0")
    if entropy > 7.2:
        flags.append("Very high entropy (packed/encrypted) — common obfuscation in malware")
        risk_score += 40
    elif entropy > 6.5:
        flags.append("Elevated entropy — file may be compressed or partially obfuscated")
        risk_score += 15

    # 3. Suspicious strings
    strings     = extract_printable_strings(data)
    found_strs  = [s for s in strings if any(kw.lower() in s.lower() for kw in SUSPICIOUS_EXE_STRINGS)]
    if found_strs:
        sample = found_strs[:5]
        flags.append(f"Suspicious strings found: {', '.join(sample[:3])}")
        risk_score += min(len(found_strs) * 5, 30)

    # 4. PE header analysis (if pefile available)
    if PEFILE_AVAILABLE:
        try:
            pe = pefile.PE(data=data)

            # Compilation timestamp
            ts = pe.FILE_HEADER.TimeDateStamp
            if ts == 0 or ts > 2_000_000_000:
                flags.append("Invalid or tampered compilation timestamp — common in malware")
                risk_score += 10

            # Suspicious imports
            if hasattr(pe, 'DIRECTORY_ENTRY_IMPORT'):
                for entry in pe.DIRECTORY_ENTRY_IMPORT:
                    for imp in entry.imports:
                        if imp.name:
                            imp_name = imp.name.decode(errors="ignore")
                            if imp_name in SUSPICIOUS_IMPORTS:
                                flags.append(f"Imports dangerous API: {imp_name}")
                                risk_score += 15

            # Section names
            for section in pe.sections:
                name = section.Name.decode(errors="ignore").strip("\x00")
                if name and not re.match(r'^[\.\w]+$', name):
                    flags.append(f"Unusual PE section name: '{name}' — seen in packed malware")
                    risk_score += 10
        except Exception as pe_err:
            flags.append(f"PE header could not be parsed: {pe_err} — may indicate a corrupted or packed file")
            risk_score += 10

    is_suspicious = risk_score >= 30
    return {
        "type":         "exe",
        "filename":     filename,
        "is_suspicious": is_suspicious,
        "risk_score":   min(risk_score, 100),
        "entropy":      entropy,
        "flags":        flags,
        "summary":      f"EXE static analysis: risk score {min(risk_score,100)}/100, entropy {entropy}/8.0"
    }


def analyze_zip(data: bytes, filename: str) -> dict:
    """
    Analyze ZIP without extracting. Safe read-only inspection.
    If inner files are PDF/doc/html/image, analyze them too (1 level deep).
    """
    flags      = []
    risk_score = 0
    inner_results = []

    # 1. Magic byte check
    if not verify_magic_bytes(data, "zip"):
        flags.append("File claims to be ZIP but internal bytes don't match — possible disguised file")
        risk_score += 25

    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return {
            "type": "zip", "filename": filename,
            "is_suspicious": True, "risk_score": 50,
            "flags": ["File is not a valid ZIP archive — could be a disguised executable"],
            "summary": "Invalid ZIP file"
        }

    # 2. Password protection
    needs_pwd = False
    for info in zf.infolist():
        if info.flag_bits & 0x1:
            needs_pwd = True
            break
    if needs_pwd:
        flags.append("ZIP is password-protected — prevents security scanning, common in phishing attachments")
        risk_score += 35

    # 3. Nesting depth (zip-in-zip)
    inner_zips = [f for f in zf.namelist() if f.lower().endswith(".zip")]
    if inner_zips:
        flags.append(f"ZIP contains nested ZIP files ({len(inner_zips)}) — evasion technique")
        risk_score += 20

    # 4. Inspect all filenames
    for name in zf.namelist():
        ext = os.path.splitext(name)[1].lower()

        # Double extension
        parts = name.lower().split(".")
        if len(parts) >= 3:
            flags.append(f"Double extension detected: '{name}' — classic trick to hide true file type")
            risk_score += 25

        # Dangerous inner extensions
        if ext in INSTANT_DANGER_EXTENSIONS:
            flags.append(f"Contains dangerous file type inside: '{name}'")
            risk_score += 40

        if ext in MACRO_EXTENSIONS:
            flags.append(f"Contains macro-enabled Office file: '{name}'")
            risk_score += 30

        if ext == ".exe":
            flags.append(f"Contains executable inside ZIP: '{name}'")
            risk_score += 40

    # 5. Compression ratio (zip bomb check)
    total_compressed   = sum(i.compress_size for i in zf.infolist())
    total_uncompressed = sum(i.file_size     for i in zf.infolist())
    if total_compressed > 0:
        ratio = total_uncompressed / total_compressed
        if ratio > 100:
            flags.append(f"Extreme compression ratio ({int(ratio)}:1) — possible zip bomb")
            risk_score += 50
        elif ratio > 20:
            flags.append(f"High compression ratio ({int(ratio)}:1) — worth noting")
            risk_score += 10

    # 6. Safe inner-file analysis (1-level deep, non-executable types only)
    if not needs_pwd:
        SAFE_INNER_TYPES = {".pdf", ".docx", ".xlsx", ".html", ".htm", ".png", ".jpg", ".jpeg"}
        for info in zf.infolist():
            ext = os.path.splitext(info.filename)[1].lower()
            if ext in SAFE_INNER_TYPES and info.file_size < 5 * 1024 * 1024:  # max 5MB inner
                try:
                    inner_data   = zf.read(info.filename)
                    inner_result = _route_attachment(inner_data, info.filename, depth=1)
                    if inner_result and inner_result.get("is_suspicious"):
                        inner_results.append(inner_result)
                        flags.append(f"Suspicious inner file: '{info.filename}' — {inner_result.get('summary','')}")
                        risk_score += 20
                except Exception:
                    pass

    is_suspicious = risk_score >= 30
    result = {
        "type":          "zip",
        "filename":      filename,
        "is_suspicious": is_suspicious,
        "risk_score":    min(risk_score, 100),
        "flags":         flags,
        "inner_results": inner_results,
        "summary":       f"ZIP analysis: {len(zf.namelist())} files inside, risk score {min(risk_score,100)}/100"
    }
    zf.close()
    return result


def analyze_image_attachment(data: bytes, filename: str) -> dict:
    """
    Analyze image attachment: magic bytes, OCR, EXIF, size ratio.
    """
    flags      = []
    risk_score = 0
    ocr_text   = ""

    # 1. Magic byte check
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if not verify_magic_bytes(data, ext):
        flags.append(f"File extension '{ext}' doesn't match internal file signature — possible disguised file")
        risk_score += 25

    if not OCR_AVAILABLE:
        return {
            "type": "image", "filename": filename,
            "is_suspicious": False, "risk_score": 0,
            "flags": ["OCR not available — image content could not be analyzed"],
            "summary": "Image analysis skipped (OCR unavailable)"
        }

    # 2. OCR text extraction
    try:
        img      = Image.open(io.BytesIO(data))
        ocr_text = pytesseract.image_to_string(img).strip()

        # 2a. Run text model on OCR result
        if ocr_text:
            text_result = run_text_model(ocr_text)
            if text_result["is_phishing"]:
                flags.append(f"OCR text classified as phishing ({text_result['confidence']}% confidence)")
                risk_score += 40

            # 2b. Check urgency keywords in OCR
            urgency_words = ["verify", "urgent", "suspended", "click here", "confirm", "account", "login", "password"]
            found_urgency = [w for w in urgency_words if w in ocr_text.lower()]
            if found_urgency:
                flags.append(f"Urgency/phishing keywords in image text: {', '.join(found_urgency)}")
                risk_score += 20

            # 2c. Extract URLs from OCR text
            ocr_urls = re.findall(r'https?://[^\s]+', ocr_text)
            for u in ocr_urls[:3]:
                if run_url_model(u):
                    flags.append(f"Phishing URL found in image text: {u}")
                    risk_score += 30

        # 3. File size vs dimension ratio (steganography hint)
        width, height = img.size
        expected_bytes = width * height * 3   # rough RGB estimate
        actual_bytes   = len(data)
        ratio = actual_bytes / max(expected_bytes, 1)
        if ratio > 3.0:
            flags.append(f"File size ({actual_bytes} bytes) is much larger than expected for its dimensions — possible hidden data")
            risk_score += 15

        # 4. EXIF metadata
        try:
            exif_data = img._getexif()
            if exif_data:
                software = exif_data.get(305, "")   # Tag 305 = Software
                if software and any(k in software.lower() for k in ["phish", "hack", "kit", "exploit"]):
                    flags.append(f"Suspicious EXIF software tag: {software}")
                    risk_score += 20
        except Exception:
            pass

    except Exception as e:
        flags.append(f"Image could not be opened: {e}")

    is_suspicious = risk_score >= 30
    return {
        "type":          "image",
        "filename":      filename,
        "is_suspicious": is_suspicious,
        "risk_score":    min(risk_score, 100),
        "flags":         flags,
        "ocr_preview":   ocr_text[:300] if ocr_text else "",
        "summary":       f"Image analysis: risk score {min(risk_score,100)}/100" + (f", OCR found {len(ocr_text)} chars" if ocr_text else "")
    }


def analyze_pdf_doc(data: bytes, filename: str) -> dict:
    """
    Analyze PDF, DOCX, XLSX files.
    Extracts text, URLs, checks for embedded JS (PDF), macros (Office), remote templates.
    """
    flags      = []
    risk_score = 0
    ext        = os.path.splitext(filename)[1].lower()
    all_text   = ""
    all_urls   = []

    # ── PDF ──────────────────────────────────────────────────────
    if ext == ".pdf":
        # 1. Magic bytes
        if not verify_magic_bytes(data, "pdf"):
            flags.append("File claims to be PDF but internal bytes don't match")
            risk_score += 25

        # 2. Embedded JavaScript
        if b"/JS" in data or b"/JavaScript" in data:
            flags.append("PDF contains embedded JavaScript — extremely rare in legitimate PDFs, common in exploits")
            risk_score += 50

        # 3. Embedded files
        if b"/EmbeddedFile" in data:
            flags.append("PDF has files embedded inside it — possible hidden executable")
            risk_score += 35

        # 4. Text + URL extraction
        if PDF_AVAILABLE:
            try:
                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    for page in pdf.pages:
                        pg_text = page.extract_text() or ""
                        all_text += pg_text
                        # Extract hyperlinks from annotations
                        for annot in (page.annots or []):
                            uri = annot.get("uri", "")
                            if uri:
                                all_urls.append(uri)
            except Exception as e:
                flags.append(f"PDF text extraction error: {e}")
        else:
            # Fallback: raw string extraction
            strings = extract_printable_strings(data)
            all_text = " ".join(strings[:200])

    # ── DOCX ─────────────────────────────────────────────────────
    elif ext == ".docx":
        if DOCX_AVAILABLE:
            try:
                doc = python_docx.Document(io.BytesIO(data))
                all_text = "\n".join(p.text for p in doc.paragraphs)
                # Extract hyperlinks
                for rel in doc.part.rels.values():
                    if "hyperlink" in rel.reltype:
                        all_urls.append(rel._target)
            except Exception as e:
                flags.append(f"DOCX extraction error: {e}")

        # Macro check via oletools
        if OLETOOLS_AVAILABLE:
            try:
                vba = VBA_Parser(filename, data=data)
                if vba.detect_vba_macros():
                    flags.append("DOCX contains VBA macros — macros can execute malicious code automatically")
                    risk_score += 35
                    for (_, _, _, vba_code) in vba.extract_macros():
                        for kw in SUSPICIOUS_MACRO_KEYWORDS:
                            if kw.lower() in vba_code.lower():
                                flags.append(f"Macro contains dangerous keyword: '{kw}'")
                                risk_score += 15
                                break
            except Exception:
                pass

        # Remote template injection check
        # DOCX is a ZIP internally — check settings.xml.rels for external URLs
        try:
            zf = zipfile.ZipFile(io.BytesIO(data))
            for name in zf.namelist():
                if "settings" in name.lower() and name.endswith(".rels"):
                    content = zf.read(name).decode(errors="ignore")
                    ext_urls = re.findall(r'Target="(https?://[^"]+)"', content)
                    if ext_urls:
                        flags.append(f"Remote template injection detected — doc fetches from: {ext_urls[0]}")
                        risk_score += 45
            zf.close()
        except Exception:
            pass

    # ── XLSX ─────────────────────────────────────────────────────
    elif ext == ".xlsx":
        if XLSX_AVAILABLE:
            try:
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(max_row=50):
                        for cell in row:
                            if cell.value and isinstance(cell.value, str):
                                all_text += cell.value + " "
                wb.close()
            except Exception as e:
                flags.append(f"XLSX extraction error: {e}")

        # Macro check for XLSX too
        if OLETOOLS_AVAILABLE:
            try:
                vba = VBA_Parser(filename, data=data)
                if vba.detect_vba_macros():
                    flags.append("XLSX contains macros — unexpected for a standard spreadsheet")
                    risk_score += 30
            except Exception:
                pass

    # ── XLSM / DOCM (macro-enabled) ──────────────────────────────
    elif ext in MACRO_EXTENSIONS:
        flags.append(f"File is macro-enabled format ({ext}) — these can auto-execute code on open")
        risk_score += 30
        if OLETOOLS_AVAILABLE:
            try:
                vba = VBA_Parser(filename, data=data)
                if vba.detect_vba_macros():
                    for (_, _, _, vba_code) in vba.extract_macros():
                        for kw in SUSPICIOUS_MACRO_KEYWORDS:
                            if kw.lower() in vba_code.lower():
                                flags.append(f"Macro contains dangerous call: '{kw}'")
                                risk_score += 15

                    # Check for auto-execution triggers
                    for trigger in ["Auto_Open", "Document_Open", "Workbook_Open", "AutoOpen", "AutoExec"]:
                        if trigger.lower() in vba_code.lower():
                            flags.append(f"Macro auto-runs on file open via '{trigger}' — high risk")
                            risk_score += 25
            except Exception:
                pass

    # ── Common: text model + URL model on extracted content ──────
    if all_text.strip():
        text_result = run_text_model(all_text[:1000])
        if text_result["is_phishing"]:
            flags.append(f"Document text classified as phishing ({text_result['confidence']}% confidence)")
            risk_score += 30

    # Also grab inline URLs from raw text
    inline_urls = re.findall(r'https?://[^\s<>"\']+', all_text)
    all_urls    = list(set(all_urls + inline_urls))[:10]

    if len(all_urls) > 7:
        flags.append(f"Unusually high number of URLs ({len(all_urls)}) — legitimate docs rarely have this many")
        risk_score += 15

    phishing_urls = []
    for u in all_urls[:5]:
        if run_url_model(u):
            phishing_urls.append(u)
            flags.append(f"Phishing URL found in document: {u}")
            risk_score += 25

    is_suspicious = risk_score >= 30
    return {
        "type":          "document",
        "filename":      filename,
        "is_suspicious": is_suspicious,
        "risk_score":    min(risk_score, 100),
        "flags":         flags,
        "urls_found":    all_urls[:5],
        "phishing_urls": phishing_urls,
        "summary":       f"Document analysis ({ext}): risk score {min(risk_score,100)}/100, {len(all_urls)} URLs found"
    }


def analyze_html_attachment(data: bytes, filename: str) -> dict:
    """
    Analyze HTML attachment for phishing indicators.
    Checks links, forms, JS obfuscation, brand impersonation, iframes, meta-refresh.
    """
    flags      = []
    risk_score = 0

    try:
        html_text = data.decode(errors="ignore")
    except Exception:
        return {
            "type": "html", "filename": filename,
            "is_suspicious": True, "risk_score": 60,
            "flags": ["HTML file could not be decoded"],
            "summary": "HTML decode failed"
        }

    if not BS4_AVAILABLE:
        # Fallback: regex only
        urls = re.findall(r'href=["\']?(https?://[^\s"\'<>]+)', html_text)
        for u in urls[:5]:
            if run_url_model(u):
                flags.append(f"Phishing URL in HTML: {u}")
                risk_score += 30
        return {
            "type": "html", "filename": filename,
            "is_suspicious": risk_score >= 30,
            "risk_score": min(risk_score, 100),
            "flags": flags,
            "summary": "Basic HTML analysis (BeautifulSoup not available)"
        }

    soup = BeautifulSoup(html_text, "html.parser")

    # 1. All href links
    all_hrefs = []
    for tag in soup.find_all("a", href=True):
        href        = tag["href"]
        display_txt = tag.get_text(strip=True)
        all_hrefs.append(href)

        # Display text vs actual URL mismatch
        if display_txt and re.match(r'https?://', display_txt):
            display_domain = re.findall(r'https?://([^/\s]+)', display_txt)
            href_domain    = re.findall(r'https?://([^/\s]+)', href)
            if display_domain and href_domain and display_domain[0] != href_domain[0]:
                flags.append(f"Link text says '{display_domain[0]}' but actually goes to '{href_domain[0]}' — classic phishing trick")
                risk_score += 40

    # Run URL model on hrefs
    phishing_hrefs = []
    for u in all_hrefs[:8]:
        if u.startswith("http") and run_url_model(u):
            phishing_hrefs.append(u)
            flags.append(f"Phishing URL in HTML: {u}")
            risk_score += 25

    # 2. Form analysis
    for form in soup.find_all("form"):
        action = form.get("action", "")
        if action.startswith("http"):
            if run_url_model(action):
                flags.append(f"Form submits data to phishing URL: {action}")
                risk_score += 45
        # Hidden fields
        hidden = form.find_all("input", {"type": "hidden"})
        if len(hidden) > 3:
            flags.append(f"Form has {len(hidden)} hidden input fields — may be used to silently steal data")
            risk_score += 15

    # 3. JavaScript obfuscation
    for script in soup.find_all("script"):
        script_text = script.get_text()
        for marker in OBFUSCATION_MARKERS:
            if marker in script_text:
                flags.append(f"Obfuscated JavaScript detected: '{marker}' — code is deliberately hidden")
                risk_score += 35
                break

    # 4. Brand impersonation: page claims to be a brand but URLs don't match
    page_text  = soup.get_text().lower()
    known_brands = {
        "paypal":    "paypal.com",
        "amazon":    "amazon.com",
        "google":    "google.com",
        "microsoft": "microsoft.com",
        "apple":     "apple.com",
        "netflix":   "netflix.com",
        "bank of america": "bankofamerica.com",
        "hdfc":      "hdfcbank.com",
        "sbi":       "sbi.co.in",
    }
    for brand, real_domain in known_brands.items():
        if brand in page_text:
            bad_links = [h for h in all_hrefs if h.startswith("http") and real_domain not in h]
            if bad_links:
                flags.append(f"Page claims to be {brand.title()} but links don't go to {real_domain} — brand impersonation")
                risk_score += 40
                break

    # 5. iFrame injection
    for iframe in soup.find_all("iframe"):
        src = iframe.get("src", "")
        if src.startswith("http"):
            flags.append(f"External iframe detected pointing to: {src[:80]}")
            risk_score += 20

    # 6. Meta refresh redirect
    for meta in soup.find_all("meta"):
        http_equiv = meta.get("http-equiv", "").lower()
        if http_equiv == "refresh":
            content = meta.get("content", "")
            redirect_url = re.findall(r'url=(.+)', content, re.IGNORECASE)
            if redirect_url:
                flags.append(f"Instant redirect (meta refresh) to: {redirect_url[0][:80]} — hides true destination")
                risk_score += 30

    # 7. Run text model on visible page text
    visible_text = soup.get_text()[:1000]
    if visible_text:
        text_result = run_text_model(visible_text)
        if text_result["is_phishing"]:
            flags.append(f"Page content classified as phishing ({text_result['confidence']}% confidence)")
            risk_score += 25

    is_suspicious = risk_score >= 30
    return {
        "type":           "html",
        "filename":       filename,
        "is_suspicious":  is_suspicious,
        "risk_score":     min(risk_score, 100),
        "flags":          flags,
        "phishing_hrefs": phishing_hrefs[:3],
        "summary":        f"HTML analysis: risk score {min(risk_score,100)}/100, {len(all_hrefs)} links checked"
    }


def _route_attachment(data: bytes, filename: str, depth: int = 0) -> dict:
    """
    Route attachment bytes to the correct analyzer based on extension.
    depth=1 means we're already inside a ZIP — don't recurse further.
    """
    ext = os.path.splitext(filename)[1].lower()

    # Instant danger extensions — no further analysis needed
    if ext in INSTANT_DANGER_EXTENSIONS:
        return {
            "type":          "script",
            "filename":      filename,
            "is_suspicious": True,
            "risk_score":    90,
            "flags":         [f"File type '{ext}' is inherently high risk — can execute system commands directly"],
            "summary":       f"High-risk script file type: {ext}"
        }

    if ext in (".exe",):
        return analyze_exe(data, filename)

    if ext in (".zip", ".rar", ".7z", ".gz", ".tar") and depth == 0:
        return analyze_zip(data, filename)

    if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"):
        return analyze_image_attachment(data, filename)

    if ext in (".pdf", ".docx", ".xlsx") or ext in MACRO_EXTENSIONS:
        return analyze_pdf_doc(data, filename)

    if ext in (".html", ".htm"):
        return analyze_html_attachment(data, filename)

    # Unknown type — basic entropy check
    entropy = calculate_entropy(data)
    flags   = []
    risk    = 0
    if entropy > 7.2:
        flags.append(f"Unknown file type with very high entropy ({entropy}) — could be an obfuscated payload")
        risk = 40
    return {
        "type":          "unknown",
        "filename":      filename,
        "is_suspicious": risk >= 30,
        "risk_score":    risk,
        "entropy":       entropy,
        "flags":         flags,
        "summary":       f"Unknown file type ({ext}), entropy: {entropy}/8.0"
    }


def scan_attachments(service, msg_id: str, payload: dict) -> list:
    """
    Find and scan all attachments in an email payload.
    Returns list of attachment result dicts.
    Safely downloads bytes via Gmail API — never executes anything.
    """
    results = []

    def _walk(parts):
        for part in parts:
            filename    = part.get("filename", "")
            body        = part.get("body", {})
            attachment_id = body.get("attachmentId")
            mime_type   = part.get("mimeType", "")

            # Recurse into multipart
            if part.get("parts"):
                _walk(part["parts"])
                continue

            if not filename or not attachment_id:
                continue

            # Skip tiny parts (likely inline images in email signature)
            if body.get("size", 0) < 100:
                continue

            # Skip very large attachments (>15MB) — too slow for real-time scan
            if body.get("size", 0) > 15 * 1024 * 1024:
                results.append({
                    "type":          "skipped",
                    "filename":      filename,
                    "is_suspicious": False,
                    "risk_score":    0,
                    "flags":         ["File too large to scan in real-time (>15MB)"],
                    "summary":       f"Skipped large file: {filename}"
                })
                continue

            try:
                att_data = service.users().messages().attachments().get(
                    userId="me", messageId=msg_id, id=attachment_id
                ).execute()
                raw_bytes = base64.urlsafe_b64decode(att_data["data"] + "==")
                result    = _route_attachment(raw_bytes, filename)
                results.append(result)
                print(f"   📎 Scanned attachment '{filename}': suspicious={result['is_suspicious']}, score={result['risk_score']}")
            except Exception as e:
                print(f"   ⚠️ Could not download attachment '{filename}': {e}")
                results.append({
                    "type":          "error",
                    "filename":      filename,
                    "is_suspicious": False,
                    "risk_score":    0,
                    "flags":         [f"Could not download attachment for scanning: {e}"],
                    "summary":       f"Download failed: {filename}"
                })

    top_parts = payload.get("parts", [])
    if top_parts:
        _walk(top_parts)

    return results


# ════════════════════════════════════════════════════════════════
#  STATIC
# ════════════════════════════════════════════════════════════════

@app.route("/")
def index():
    return send_from_directory("static", "index.html")


# ════════════════════════════════════════════════════════════════
#  OAUTH ROUTES
# ════════════════════════════════════════════════════════════════

@app.route("/login")
def login():
    if not os.path.exists(CLIENT_SECRETS_FILE):
        return jsonify({"error": "client_secret.json not found."}), 503

    flow = Flow.from_client_secrets_file(
        CLIENT_SECRETS_FILE, scopes=SCOPES, redirect_uri=REDIRECT_URI
    )
    auth_url, state = flow.authorization_url(
        access_type="offline", include_granted_scopes="true", prompt="consent"
    )
    session["state"] = state
    return redirect(auth_url)


@app.route("/oauth2callback")
def oauth2callback():
    print("OAuth callback received")
    if "error" in request.args:
        return f"OAuth error: {request.args['error']}", 400

    flow = Flow.from_client_secrets_file(
        CLIENT_SECRETS_FILE,
        scopes=SCOPES,
        state=session.get("state"),
        redirect_uri=REDIRECT_URI,
    )
    try:
        flow.fetch_token(authorization_response=request.url)
    except Exception as e:
        print(f"❌ fetch_token failed: {e}\n{traceback.format_exc()}")
        return f"Token error: {e}", 500

    creds = flow.credentials
    session["credentials"] = {
        "token":         creds.token,
        "refresh_token": creds.refresh_token,
        "token_uri":     creds.token_uri,
        "client_id":     creds.client_id,
        "client_secret": creds.client_secret,
        "scopes":        list(creds.scopes),
    }
    print(f"✅ OAuth complete. Refresh token present: {bool(creds.refresh_token)}")
    return redirect("/?gmail=1")


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/")


# ════════════════════════════════════════════════════════════════
#  AUTH STATUS
# ════════════════════════════════════════════════════════════════

@app.route("/auth/status")
def auth_status():
    creds = get_credentials()
    if not creds:
        return jsonify({"logged_in": False})
    try:
        service   = build("oauth2", "v2", credentials=creds)
        user_info = service.userinfo().get().execute()
        return jsonify({
            "logged_in": True,
            "email":     user_info.get("email"),
            "name":      user_info.get("name"),
            "picture":   user_info.get("picture"),
        })
    except Exception as e:
        return jsonify({"logged_in": False, "error": str(e)})


# ════════════════════════════════════════════════════════════════
#  GMAIL — SCAN  (now includes attachment scanning)
# ════════════════════════════════════════════════════════════════

@app.route("/gmail/scan")
def gmail_scan():
    creds = get_credentials()
    if not creds:
        return jsonify({"error": "Not authenticated"}), 401

    if text_model is None or tfidf is None:
        return jsonify({"error": "Models not loaded"}), 503

    max_emails  = min(int(request.args.get("max", 30)), 100)
    inbox_quota = max_emails // 2
    spam_quota  = max_emails - inbox_quota
    print(f"📬 Gmail scan: {inbox_quota} inbox + {spam_quota} spam")

    try:
        service = build("gmail", "v1", credentials=creds)

        trusted_contacts = get_trusted_contacts(service)

        inbox_resp = service.users().messages().list(
            userId="me", maxResults=inbox_quota, labelIds=["INBOX"]
        ).execute()
        inbox_msgs = [(m, "INBOX") for m in inbox_resp.get("messages", [])]

        spam_resp = service.users().messages().list(
            userId="me", maxResults=spam_quota, labelIds=["SPAM"]
        ).execute()
        spam_msgs = [(m, "SPAM") for m in spam_resp.get("messages", [])]

        all_messages = spam_msgs + inbox_msgs
        print(f"   Got {len(inbox_msgs)} inbox + {len(spam_msgs)} spam = {len(all_messages)} total")

    except Exception as e:
        print(f"❌ Gmail list failed: {e}\n{traceback.format_exc()}")
        return jsonify({"error": f"Gmail API error: {str(e)}"}), 500

    if not all_messages:
        return jsonify({
            "total": 0, "phishing_count": 0,
            "safe_count": 0, "phishing_rate": 0, "emails": []
        })

    emails         = []
    phishing_count = 0

    for msg_meta, folder in all_messages:
        try:
            msg_data = service.users().messages().get(
                userId="me", id=msg_meta["id"], format="full"
            ).execute()
        except Exception as e:
            print(f"   skip {msg_meta['id']}: {e}")
            continue

        headers = {
            h["name"]: h["value"]
            for h in msg_data.get("payload", {}).get("headers", [])
        }
        subject = headers.get("Subject", "(No Subject)")
        sender  = headers.get("From",    "Unknown")
        date    = headers.get("Date",    "")
        snippet = msg_data.get("snippet", "")

        body            = decode_body(msg_data.get("payload", {}))
        text_to_analyze = (body or snippet or subject)[:1000]

        # 1. TEXT CHECK
        try:
            vec        = tfidf.transform([text_to_analyze])
            pred_text  = int(text_model.predict(vec)[0])
            proba_text = text_model.predict_proba(vec)[0]
            text_conf  = round(float(max(proba_text)) * 100, 2)
        except Exception:
            pred_text = 0; text_conf = 0.0

        # 2. URL CHECK
        urls_found = re.findall(r'https?://[^\s<>"\']+', snippet + " " + body)
        urls_found = list(set(urls_found))[:5]

        has_malicious_url = False
        url_alerts = []
        if url_model is not None and url_scaler is not None:
            for u in urls_found:
                try:
                    f = extract_url_features(u)
                    scaled = url_scaler.transform([f])
                    u_pred = int(url_model.predict(scaled)[0])
                    if u_pred == 1:
                        has_malicious_url = True
                        url_alerts.append(u)
                except Exception:
                    pass

        # 3. IMPERSONATION CHECK
        impersonation_result = check_impersonation(sender, trusted_contacts)

        # 4. ATTACHMENT SCAN  ← NEW
        attachment_results   = []
        has_suspicious_attachment = False
        try:
            attachment_results = scan_attachments(service, msg_meta["id"], msg_data.get("payload", {}))
            has_suspicious_attachment = any(a.get("is_suspicious") for a in attachment_results)
        except Exception as att_err:
            print(f"   ⚠️ Attachment scan failed for {msg_meta['id']}: {att_err}")

        in_spam     = (folder == "SPAM")
        is_phishing = (
            (pred_text == 1)
            or in_spam
            or has_malicious_url
            or impersonation_result["is_impersonation"]
            or has_suspicious_attachment   # ← NEW signal
        )
        if is_phishing:
            phishing_count += 1

        emails.append({
            "id":             msg_meta["id"],
            "subject":        subject,
            "sender":         sender,
            "date":           date,
            "snippet":        snippet[:200],
            "body_preview":   text_to_analyze[:500],
            "folder":         folder,
            "is_phishing":    is_phishing,
            "urls_found":     urls_found,
            "text_result":    {"confidence": text_conf, "is_phishing": bool(pred_text == 1)},
            "url_result":     {"has_malicious_url": has_malicious_url, "malicious_urls": url_alerts},
            "impersonation":  impersonation_result,
            # ── NEW ──
            "attachment_result": {
                "has_attachments":         len(attachment_results) > 0,
                "attachment_count":        len(attachment_results),
                "has_suspicious_attachment": has_suspicious_attachment,
                "attachments":             attachment_results,
            },
        })

    safe_count    = len(emails) - phishing_count
    phishing_rate = round(phishing_count / max(len(emails), 1) * 100, 1)
    print(f"✅ Scan done: {phishing_count} phishing / {safe_count} safe")

    return jsonify({
        "total":          len(emails),
        "phishing_count": phishing_count,
        "safe_count":     safe_count,
        "phishing_rate":  phishing_rate,
        "emails":         emails,
    })


# ════════════════════════════════════════════════════════════════
#  EXPLAIN — EMAIL (Gemini powered)
# ════════════════════════════════════════════════════════════════

@app.route("/explain/email", methods=["POST"])
def explain_email():
    data    = request.get_json()
    subject = (data or {}).get("subject", "")
    sender  = (data or {}).get("sender", "")
    snippet = (data or {}).get("snippet", "")
    body    = (data or {}).get("body_preview", "")
    urls    = (data or {}).get("urls_found", [])
    conf    = (data or {}).get("confidence", 0)
    folder  = (data or {}).get("folder", "INBOX")
    # Attachment context passed from frontend
    att_summary = (data or {}).get("attachment_summary", "")

    att_context = f"\n- Attachment findings: {att_summary}" if att_summary else ""

    prompt = f"""You are a cybersecurity expert and educator helping regular users understand phishing threats.

Analyze this email and provide a clear, educational explanation for a non-technical person.

EMAIL DETAILS:
- Subject: {subject}
- Sender: {sender}
- Folder: {folder}
- AI Confidence it's phishing: {conf}%
- URLs found: {', '.join(urls) if urls else 'None'}
- Content preview: {(body or snippet)[:400]}{att_context}

Respond ONLY with a valid JSON object in this exact format (no markdown, no extra text):
{{
  "verdict_summary": "One punchy sentence explaining why this is phishing",
  "red_flags": [
    "Specific red flag 1 found in this email",
    "Specific red flag 2 found in this email",
    "Specific red flag 3 found in this email"
  ],
  "what_attacker_wants": "What the attacker is trying to steal or achieve",
  "what_to_do": "Specific action the user should take right now",
  "how_to_spot_next_time": "One key lesson to remember for future emails",
  "danger_level": "HIGH" or "MEDIUM" or "LOW",
  "danger_reason": "Brief reason for the danger level"
}}"""

    ai_text = call_gemini(prompt)

    if not ai_text:
        flags = []
        if "urgent" in (subject + snippet).lower() or "immediately" in (subject + snippet).lower():
            flags.append("Uses urgency language to pressure you into acting fast")
        if urls:
            flags.append(f"Contains {len(urls)} suspicious link(s)")
        if folder == "SPAM":
            flags.append("Gmail's own filters marked this as spam")
        if conf > 80:
            flags.append(f"AI model is {conf}% confident this is phishing")
        if att_summary:
            flags.append(f"Suspicious attachment: {att_summary}")

        return jsonify({
            "verdict_summary": "This email shows multiple signs of a phishing attempt.",
            "red_flags": flags or ["Suspicious content pattern detected by AI", "Unusual sender behavior", "Content matches known phishing templates"],
            "what_attacker_wants": "Likely trying to steal your login credentials or personal information.",
            "what_to_do": "Do not click any links or open attachments. Mark as spam and delete immediately.",
            "how_to_spot_next_time": "Legitimate companies never ask for sensitive info via email urgently.",
            "danger_level": "HIGH" if conf > 80 else "MEDIUM",
            "danger_reason": "High AI confidence score with suspicious content patterns."
        })

    try:
        clean = ai_text.strip()
        if clean.startswith("```"):
            clean = re.sub(r"```[a-z]*\n?", "", clean).strip().rstrip("`").strip()
        result = json.loads(clean)
        return jsonify(result)
    except Exception as e:
        print(f"❌ Gemini JSON parse error: {e}\nRaw: {ai_text}")
        return jsonify({
            "verdict_summary": "This email shows signs of a phishing attempt.",
            "red_flags": ["Suspicious sender pattern", "Content matches phishing templates", "Unusual link structure"],
            "what_attacker_wants": "Trying to steal credentials or personal data.",
            "what_to_do": "Do not click any links or open attachments. Delete this email immediately.",
            "how_to_spot_next_time": "Always verify sender identity through official channels.",
            "danger_level": "HIGH",
            "danger_reason": "Multiple phishing indicators detected."
        })


# ════════════════════════════════════════════════════════════════
#  EXPLAIN — ATTACHMENT  (NEW — Gemini powered)
# ════════════════════════════════════════════════════════════════

@app.route("/explain/attachment", methods=["POST"])
def explain_attachment():
    """
    Takes a single attachment result dict and asks Gemini to explain
    in plain English why it's suspicious.
    """
    data       = request.get_json() or {}
    filename   = data.get("filename", "unknown")
    file_type  = data.get("type", "unknown")
    risk_score = data.get("risk_score", 0)
    flags      = data.get("flags", [])
    summary    = data.get("summary", "")
    entropy    = data.get("entropy", None)

    entropy_line = f"\n- Entropy: {entropy}/8.0 (7.2+ = packed/obfuscated)" if entropy is not None else ""

    prompt = f"""You are a cybersecurity educator explaining a suspicious email attachment to an everyday user.

ATTACHMENT DETAILS:
- Filename: {filename}
- File type: {file_type}
- Risk score: {risk_score}/100
- Summary: {summary}{entropy_line}
- Detected issues:
{chr(10).join(f'  * {f}' for f in flags[:8])}

Explain this clearly to a non-technical person. Respond ONLY with valid JSON (no markdown):
{{
  "verdict_summary": "Plain English: why is this file dangerous?",
  "red_flags": ["Issue 1 explained simply", "Issue 2 explained simply", "Issue 3 explained simply"],
  "what_happens_if_opened": "Exactly what could happen if the user opens this file",
  "what_to_do": "What the user should do right now",
  "technical_insight": "One sentence explaining the attack technique used",
  "danger_level": "HIGH" or "MEDIUM" or "LOW"
}}"""

    ai_text = call_gemini(prompt)

    if not ai_text:
        return jsonify({
            "verdict_summary": f"This {file_type} file has a risk score of {risk_score}/100 and shows signs of being malicious.",
            "red_flags": flags[:3] or ["Suspicious file characteristics detected"],
            "what_happens_if_opened": "Could install malware, steal your data, or give attackers access to your computer.",
            "what_to_do": "Do NOT open this file. Delete the email immediately and report it to your IT team.",
            "technical_insight": summary,
            "danger_level": "HIGH" if risk_score >= 60 else "MEDIUM" if risk_score >= 30 else "LOW"
        })

    try:
        clean = ai_text.strip()
        if clean.startswith("```"):
            clean = re.sub(r"```[a-z]*\n?", "", clean).strip().rstrip("`").strip()
        return jsonify(json.loads(clean))
    except Exception as e:
        print(f"❌ Gemini JSON parse error: {e}")
        return jsonify({
            "verdict_summary": f"File flagged with risk score {risk_score}/100.",
            "red_flags": flags[:3] or ["Suspicious characteristics"],
            "what_happens_if_opened": "May execute malicious code or steal data.",
            "what_to_do": "Do not open this file. Delete and report.",
            "technical_insight": summary,
            "danger_level": "HIGH" if risk_score >= 60 else "MEDIUM"
        })


# ════════════════════════════════════════════════════════════════
#  EXPLAIN — URL (Gemini powered)
# ════════════════════════════════════════════════════════════════

@app.route("/explain/url", methods=["POST"])
def explain_url():
    data       = request.get_json()
    url        = (data or {}).get("url", "")
    prediction = (data or {}).get("prediction", "phishing")
    confidence = (data or {}).get("confidence", 0)

    features   = extract_url_features(url)
    red_flags  = get_url_red_flags(url)

    feature_desc = f"""URL Analysis Features:
- Length: {features[0]} chars
- Dots: {features[1]}
- Hyphens: {features[2]}
- Slashes: {features[3]}
- @ symbols: {features[4]}
- Has IP address: {'Yes' if features[7] else 'No'}
- Uses HTTPS: {'Yes' if features[8] else 'No'}
- Digit ratio: {features[9]}"""

    prompt = f"""You are a cybersecurity educator explaining phishing URLs to everyday users.

URL: {url}
AI Verdict: {prediction.upper()} ({confidence}% confidence)
{feature_desc}
Pre-detected red flags: {', '.join(red_flags) if red_flags else 'None obvious'}

Respond ONLY with a valid JSON object (no markdown):
{{
  "verdict_summary": "Plain English explanation of why this URL is {'dangerous' if prediction == 'phishing' else 'safe'}",
  "red_flags": ["Specific red flag 1", "Specific red flag 2", "Specific red flag 3"],
  "what_happens_if_clicked": "Exactly what could happen if someone visits this URL",
  "how_to_verify": "How a user can verify if a URL is safe before clicking",
  "safe_alternative": "What the legitimate version of this URL might look like (if phishing)",
  "danger_level": "HIGH" or "MEDIUM" or "LOW",
  "tip_for_future": "One memorable tip to identify similar URLs"
}}"""

    ai_text = call_gemini(prompt)

    if not ai_text:
        return jsonify({
            "verdict_summary": f"This URL appears {'dangerous' if prediction == 'phishing' else 'safe'} with {confidence}% confidence.",
            "red_flags": red_flags or ["No obvious red flags detected"],
            "what_happens_if_clicked": "Could lead to credential theft or malware installation.",
            "how_to_verify": "Hover over the link to preview the URL. Use Google Safe Browsing to check.",
            "safe_alternative": "Always type URLs directly into your browser instead of clicking links.",
            "danger_level": "HIGH" if confidence > 80 else "MEDIUM",
            "tip_for_future": "Real banks and services never use IP addresses or excessive hyphens in URLs."
        })

    try:
        clean = ai_text.strip()
        if clean.startswith("```"):
            clean = re.sub(r"```[a-z]*\n?", "", clean).strip().rstrip("`").strip()
        return jsonify(json.loads(clean))
    except Exception as e:
        print(f"❌ Gemini JSON parse error: {e}")
        return jsonify({
            "verdict_summary": f"URL classified as {prediction} with {confidence}% confidence.",
            "red_flags": red_flags or ["Suspicious URL pattern"],
            "what_happens_if_clicked": "May lead to credential theft.",
            "how_to_verify": "Check with Google Safe Browsing.",
            "safe_alternative": "Navigate directly to official websites.",
            "danger_level": "HIGH" if confidence > 80 else "MEDIUM",
            "tip_for_future": "Look for HTTPS and recognize the official domain."
        })


# ════════════════════════════════════════════════════════════════
#  EXPLAIN — TEXT / OCR (Gemini powered)
# ════════════════════════════════════════════════════════════════

@app.route("/explain/text", methods=["POST"])
def explain_text():
    data       = request.get_json()
    text       = (data or {}).get("text", "")
    prediction = (data or {}).get("prediction", "phishing")
    confidence = (data or {}).get("confidence", 0)
    source     = (data or {}).get("source", "email/sms")

    prompt = f"""You are a cybersecurity educator. A user submitted a {source} for phishing analysis.

Content: {text[:500]}
AI Verdict: {prediction.upper()} ({confidence}% confidence)

Respond ONLY with a valid JSON object (no markdown):
{{
  "verdict_summary": "Plain English summary of the verdict",
  "red_flags": ["Specific red flag 1 from the actual content", "Specific red flag 2", "Specific red flag 3"],
  "psychological_tricks": ["Trick 1 the attacker uses (e.g., urgency, fear, greed)", "Trick 2"],
  "what_attacker_wants": "Specific goal of this phishing attempt",
  "what_to_do": "Exactly what the user should do right now",
  "how_to_verify_legitimacy": "How to check if this message is actually from a real organization",
  "danger_level": "HIGH" or "MEDIUM" or "LOW",
  "educational_insight": "One important cybersecurity lesson from this example"
}}"""

    ai_text = call_gemini(prompt)

    if not ai_text:
        return jsonify({
            "verdict_summary": f"This {source} appears {'dangerous' if prediction == 'phishing' else 'safe'}.",
            "red_flags": ["Suspicious language pattern", "Pressure tactics detected", "Unusual request"],
            "psychological_tricks": ["Creates urgency to prevent careful thinking", "Impersonates trusted authority"],
            "what_attacker_wants": "Steal personal information or account credentials.",
            "what_to_do": "Do not respond or click any links. Report and delete.",
            "how_to_verify_legitimacy": "Contact the organization directly through their official website.",
            "danger_level": "HIGH" if confidence > 80 else "MEDIUM",
            "educational_insight": "Legitimate organizations never request sensitive data via unsolicited messages."
        })

    try:
        clean = ai_text.strip()
        if clean.startswith("```"):
            clean = re.sub(r"```[a-z]*\n?", "", clean).strip().rstrip("`").strip()
        return jsonify(json.loads(clean))
    except Exception as e:
        print(f"❌ Gemini JSON parse error: {e}")
        return jsonify({
            "verdict_summary": f"Content classified as {prediction}.",
            "red_flags": ["Suspicious content detected"],
            "psychological_tricks": ["Urgency manipulation"],
            "what_attacker_wants": "Personal information theft.",
            "what_to_do": "Delete and report.",
            "how_to_verify_legitimacy": "Contact organization via official channels.",
            "danger_level": "HIGH",
            "educational_insight": "Always verify unexpected urgent requests."
        })


# ════════════════════════════════════════════════════════════════
#  MANUAL — URL
# ════════════════════════════════════════════════════════════════

@app.route("/analyze/url", methods=["POST"])
def analyze_url():
    if url_model is None or url_scaler is None:
        return jsonify({"error": "URL model not loaded"}), 503

    data = request.get_json()
    url  = (data or {}).get("url", "").strip()
    if not url:
        return jsonify({"error": "No URL provided"}), 400

    try:
        features  = extract_url_features(url)
        scaled    = url_scaler.transform([features])
        pred      = int(url_model.predict(scaled)[0])
        proba     = url_model.predict_proba(scaled)[0]
        red_flags = get_url_red_flags(url)
        return jsonify({
            "url":        url,
            "prediction": "phishing" if pred == 1 else "safe",
            "confidence": round(float(max(proba)) * 100, 2),
            "red_flags":  red_flags,
            "features": {
                "length":      features[0],
                "dots":        features[1],
                "hyphens":     features[2],
                "has_ip":      bool(features[7]),
                "uses_https":  bool(features[8]),
                "digit_ratio": features[9],
            }
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ════════════════════════════════════════════════════════════════
#  MANUAL — TEXT
# ════════════════════════════════════════════════════════════════

@app.route("/analyze/text", methods=["POST"])
def analyze_text():
    if text_model is None or tfidf is None:
        return jsonify({"error": "Text model not loaded"}), 503

    data = request.get_json()
    text = (data or {}).get("text", "").strip()
    if not text:
        return jsonify({"error": "No text provided"}), 400

    try:
        vec   = tfidf.transform([text])
        pred  = int(text_model.predict(vec)[0])
        proba = text_model.predict_proba(vec)[0]
        return jsonify({
            "prediction": "phishing" if pred == 1 else "safe",
            "confidence": round(float(max(proba)) * 100, 2),
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ════════════════════════════════════════════════════════════════
#  SCREENSHOT / OCR
# ════════════════════════════════════════════════════════════════

@app.route("/predict/screenshot", methods=["POST"])
def predict_screenshot():
    if not OCR_AVAILABLE:
        return jsonify({"error": "OCR not available — install Pillow and pytesseract"}), 503
    if text_model is None or tfidf is None:
        return jsonify({"error": "Text model not loaded"}), 503

    if "image" not in request.files:
        return jsonify({"error": "No image uploaded"}), 400

    try:
        file           = request.files["image"]
        image          = Image.open(io.BytesIO(file.read()))
        extracted_text = pytesseract.image_to_string(image).strip()
    except Exception as e:
        return jsonify({"error": f"OCR error: {str(e)}"}), 500

    if not extracted_text:
        return jsonify({"error": "No text found in image"}), 422

    try:
        vec   = tfidf.transform([extracted_text])
        pred  = int(text_model.predict(vec)[0])
        proba = text_model.predict_proba(vec)[0]
        return jsonify({
            "prediction": "phishing" if pred == 1 else "safe",
            "confidence": round(float(max(proba)) * 100, 2),
            "ocr_text":   extracted_text[:500],
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ════════════════════════════════════════════════════════════════
#  RUN
# ════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    print("\n🛡️  PhishGuard starting on http://localhost:5000")
    app.run(debug=True, port=5000, use_reloader=False)